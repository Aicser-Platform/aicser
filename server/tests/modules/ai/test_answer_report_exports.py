"""Any answer as a report: one document for the page, the PDF and every export, charts included."""

import asyncio
import io
import struct
import zlib

import pytest

from ee.modules.ai.reports.operations import answer_report
from ee.modules.ai.services import export_artifacts_service as X


def _png(w=40, h=20):
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


BAR = {"title": {"text": "Sales by region"}, "xAxis": {"data": ["N", "S"]}, "series": [{"type": "bar", "data": [3, 5]}]}
PIE = {"title": {"text": "Share"}, "series": [{"type": "pie", "data": [{"name": "N", "value": 3}]}], "_chart_query": {}}

META = {
    "executive_summary": "South leads.",
    "echarts_config": BAR,
    "execution_metadata": {"complementary_charts": [BAR, PIE]},
    "insights": [{"title": "South leads", "description": "5 vs 3"}, "North is flat"],
    "recommendations": [{"action": "Stock South"}, "Review North"],
    "query_result": [{"region": "N", "sales": 3}, {"region": "S", "sales": 5}],
    "query_result_count": 2,
    "sql_query": "SELECT region, SUM(amount) FROM orders GROUP BY 1",
}


def test_an_ordinary_answer_becomes_a_report_with_every_chart():
    r = answer_report(META, question="Sales by region?", answer="**South** sells more.", generated_at="2026")
    charts = [s for s in r["sections"] if s["type"] == "chart"]
    assert [c["title"] for c in charts] == ["Sales by region", "Share"]  # main + More views, no duplicate
    assert "_chart_query" not in charts[1]["chart"]
    assert charts[0]["narrative"] == "**South** sells more."
    kinds = [s["id"] for s in r["sections"]]
    assert "answer-findings" in kinds and "answer-data" in kinds
    assert r["recommendations"] == [{"action": "Stock South", "title": "Stock South"}, {"title": "Review North"}]
    assert r["provenance"]["sql"].startswith("SELECT") and r["title"] == "Sales by region?"


def test_nothing_to_show_is_not_a_report():
    assert answer_report({}, question="hi", answer="") is None


def test_chart_options_follow_reading_order_without_duplicates_or_scorecards():
    ctx = {**META, "report_sections": [{"title": "KPIs", "chart": {"type": "scorecard"}}, {"title": "Trend", "chart": {"series": [{"type": "line", "data": [1]}]}}]}
    assert [c["title"] for c in X.chart_options(ctx)] == ["Sales by region", "Share", "Trend"]


def test_markdown_becomes_word_formatting():
    from docx import Document

    doc = Document()
    X._add_markdown(doc, "## Findings\n- **South** leads\n1. Act *now*\nPlain `code`")
    texts = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Heading 3", "Findings") in texts
    assert ("List Bullet", "South leads") in texts and ("List Number", "Act now") in texts
    bullet = next(p for p in doc.paragraphs if p.text == "South leads")
    assert bullet.runs[0].bold is True


@pytest.fixture
def fake_charts(monkeypatch, tmp_path):
    monkeypatch.setattr(X, "_ARTIFACTS_ROOT", tmp_path)

    async def images(ctx):
        return [{"title": c["title"], "png": _png()} for c in X.chart_options(ctx)]

    async def no_plan(ctx, audience):
        return None

    async def allowed(ctx, fmt):
        return None

    monkeypatch.setattr(X, "chart_images", images)
    monkeypatch.setattr(X, "_llm_plan_report", no_plan)
    import ee.modules.ai.services.skill_entitlements as ent
    monkeypatch.setattr(ent, "check_export_format_entitlement", allowed)


def test_word_export_contains_the_charts(fake_charts):
    from docx import Document

    out = asyncio.run(X.generate_docx({**META, "query": "Sales by region?", "organization_id": "o"}))
    assert out["success"]
    doc = Document(out["path"])
    assert len(doc.inline_shapes) == 2
    assert any(p.text.startswith("Figure 1. Sales by region") for p in doc.paragraphs)


def test_plain_deck_has_an_exhibit_slide_per_chart(fake_charts):
    from pptx import Presentation

    out = asyncio.run(X._generate_pptx_plain({**META, "query": "Sales by region?", "organization_id": "o"}))
    prs = Presentation(out["path"])
    pictures = [sh for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13]
    assert len(pictures) == 2


def test_picture_fits_its_box_without_distortion():
    from pptx import Presentation

    from ee.modules.ai.services.presentation_templates import add_picture_fit

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = add_picture_fit(slide, _png(200, 100), 0, 0, 4, 4)
    assert abs(pic.width / pic.height - 2.0) < 0.01 and pic.width <= 4 * 914400 + 1
