"""Decks planned from the request, held to the evidence, drawn in the brand."""

import asyncio
import struct
import zlib

from ee.modules.ai.services import deck_builder as B
from ee.modules.ai.services import deck_planner as P


def _png(w=40, h=20):
    raw = b"".join(b"\x00" + b"\x10\x80\xc0" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


CHARTS = [
    {"title": "Revenue by segment", "option": {"series": [{"type": "bar", "data": [4612, 1200, 800]}], "xAxis": {"data": ["VIP", "Mid", "Low"]}}},
    {"title": "Churn by tenure", "option": {"series": [{"type": "line", "data": [0.31, 0.12]}]}},
    {"title": "Unused chart", "option": {"series": [{"type": "pie", "data": [{"value": 5}]}]}},
]
CTX = {
    "query": "Customer analytics with actionable recommendations as a 20 page deck for leadership",
    "executive_summary": "VIP customers are 10% of customers and drive 46% of revenue. Churn is 31% in the first year.",
    "insights": [{"title": "VIP concentration", "description": "Top 10% drive 46% of revenue"}],
    "recommendations": [{"action": "Launch a VIP retention programme", "impact": "protect 46% of revenue"}],
    "query_result": [{"segment": "VIP", "revenue": 4612.4}],
}


def test_figures_must_come_from_the_analysis():
    facts = P.gather_facts(CTX, CHARTS)
    assert P._supported(46, facts["numbers"]) and P._supported(0.31 * 100, facts["numbers"])
    assert P._supported(4.6e3, facts["numbers"])        # $4.6K ≈ 4,612
    assert not P._supported(73, facts["numbers"])       # invented
    assert P._claimed_numbers("3 actions for 2026") == []  # counts and years aren't data claims


def test_plan_is_checked_before_drawing():
    raw = {
        "title": "Customer analytics", "requested_slides": 20,
        "slides": [
            {"layout": "executive_summary", "title": "VIPs drive 46% of revenue", "bullets": ["VIPs drive 46% of revenue", "Margin grew 73%"]},
            {"layout": "chart", "title": "Churn hits 31% in year one", "chart": 1, "bullets": ["31% churn in year one"]},
            {"layout": "chart", "title": "Made-up 88% uplift", "chart": 0, "bullets": []},
            {"layout": "chart", "title": "Ghost chart", "chart": 9},
            {"layout": "recommendations", "title": "Protect the VIP base", "items": [{"action": "Launch a VIP retention programme", "priority": "high"}, {"action": "Cut prices 50%"}]},
        ],
    }
    plan = P.check_plan(raw, P.gather_facts(CTX, CHARTS), CHARTS, CTX)
    s = plan["slides"]
    assert s[0]["layout"] == "cover"
    assert s[1]["bullets"] == ["VIPs drive 46% of revenue"]           # 73% dropped
    assert s[3]["title"] == "Revenue by segment"                       # unsupported 88% → topic
    assert not any(x.get("title") == "Ghost chart" for x in s)         # no chart, no content → gone
    assert [i["action"] for i in s[4]["items"]] == ["Launch a VIP retention programme"]
    assert any(x.get("exhibit") and x["chart"] == 2 for x in s)        # unused chart kept as exhibit
    assert s[-1]["layout"] == "appendix"
    assert plan["check"]["requested_slides"] == 20 and "20" in plan["check"]["shortfall"]
    assert plan["check"]["bullets_removed"] == 2 and plan["check"]["titles_softened"] == 1


def test_requested_length_only_raises_report_depth():
    assert P.tier_for_length(20, "standard") == "long"
    assert P.tier_for_length(10, "brief") == "standard"
    assert P.tier_for_length(5, "long") == "long"
    assert P.tier_for_length(None, "standard") == "standard"


def test_deck_is_drawn_in_the_brand_with_notes(tmp_path):
    from pptx import Presentation

    from ee.modules.ai.services.brand_config_service import BrandConfig

    plan = P.check_plan({"title": "Customer analytics", "subtitle": "Q3 review", "audience": "Leadership", "slides": [
        {"layout": "cover", "title": "Customer analytics"},
        {"layout": "kpis", "title": "Where we stand", "bullets": ["VIPs drive 46% of revenue"]},
        {"layout": "chart", "title": "VIPs drive 46% of revenue", "chart": 0, "bullets": ["Top 10% drive 46%"], "notes": "Say this."},
        {"layout": "recommendations", "title": "Protect the VIP base", "items": [{"action": "Launch a VIP retention programme", "priority": "high"}]},
        {"layout": "closing", "title": "Thank you"},
    ]}, P.gather_facts(CTX, CHARTS), CHARTS, CTX)
    brand = BrandConfig(primary="#ff6600", secondary="#112233", font_family="Georgia", org_name="Acme")
    out = B.build_planned_deck(plan, [_png(), _png(), _png()], brand, CTX, tmp_path / "d.pptx", logo=_png(60, 20))
    prs = Presentation(str(out))
    assert len(prs.slides) == len(plan["slides"])
    fonts = {r.font.name for sl in prs.slides for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert "Georgia" in fonts
    pictures = sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13)
    assert pictures >= 4  # charts + logo on cover and footers
    notes = [sl.notes_slide.notes_text_frame.text for sl in prs.slides if sl.has_notes_slide]
    assert "Say this." in notes


def test_logo_fetch_refuses_private_addresses_and_accepts_inline_images():
    assert asyncio.run(B.fetch_logo("http://127.0.0.1/logo.png")) is None
    assert asyncio.run(B.fetch_logo("http://169.254.169.254/latest/meta-data")) is None
    assert asyncio.run(B.fetch_logo("file:///etc/passwd")) is None
    import base64
    data = "data:image/png;base64," + base64.b64encode(_png()).decode()
    assert asyncio.run(B.fetch_logo(data)) == _png()


def test_brand_palette_goes_onto_charts():
    opt = B.brand_chart_option({"series": []}, ["#ff6600", "#112233"], "Georgia")
    assert opt["color"] == ["#ff6600", "#112233"] and opt["textStyle"]["fontFamily"] == "Georgia"


def test_deliverable_runs_the_analysis_first_only_when_there_is_none(monkeypatch):
    import ee.modules.ai.kernel.capability_registry as reg
    from ee.modules.ai.nodes import skill_executor_node as X

    calls = []

    async def fake_report(ctx):
        calls.append(ctx["workflow_state"]["query"])
        ws = ctx["workflow_state"]
        ws.update({"report_sections": [{"title": "Churn", "status": "complete"}], "executive_summary": "done"})
        return {"success": True, "workflow_state": ws}

    async def quiet(*a, **k):
        return None

    monkeypatch.setattr(reg, "_exec_executive_report", fake_report)
    monkeypatch.setattr(X, "_emit_skill_step_event", quiet)
    services = {"multi_query_service": object(), "data_service": object()}

    state, ctx = {"query": "20 page deck", "data_source_id": "ds"}, dict(services)
    asyncio.run(X._ensure_analysis_for_deliverable(state, ctx, None))
    assert calls == ["20 page deck"] and ctx["report_sections"] and state["executive_summary"] == "done"

    asyncio.run(X._ensure_analysis_for_deliverable({"query": "q", "data_source_id": "ds"}, {**services, "query_result": [{"a": 1}]}, None))
    asyncio.run(X._ensure_analysis_for_deliverable({"query": "q"}, dict(services), None))  # no data source
    assert len(calls) == 1


def test_a_deck_without_data_states_no_figures():
    plan = P.check_plan({"title": "Retention strategy", "slides": [
        {"layout": "bullets", "title": "Churn costs firms 25% of revenue", "bullets": ["Segment customers by value", "Most firms lose 30% a year"]},
    ]}, P.gather_facts({}, []), [], {"query": "a retention strategy deck"})
    content = next(s for s in plan["slides"] if s["layout"] == "bullets")
    assert content["bullets"] == ["Segment customers by value"]
    assert "25" not in content["title"]
